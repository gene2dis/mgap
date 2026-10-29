"""Parser and manifest tests for bin/mgap_report.py.

The files in data/ are excerpts of real MGAP outputs (Phase 8a run of the Illumina routing set,
2026-10-25), with work-directory paths removed. Run with: pytest tests/bin
"""

import io
import json
import sys
from pathlib import Path

import pytest

import mgap_report as report

# --------------------------------------------------------------------------------------------------
# Manifest
# --------------------------------------------------------------------------------------------------


def test_manifest_tsv_relative_paths_and_run_level_rows(tmp_path, data):
    (tmp_path / "files").mkdir()
    (tmp_path / "files" / "a.json").write_text("{}")
    manifest = tmp_path / "manifest.tsv"
    manifest.write_text(
        "sample\ttool\tfile\n"
        "S1\tfastp\tfiles/a.json\n"
        f"-\tcitations\t{data / 'CITATIONS.md'}\n"
        "S1\tunknown_tool\tfiles/a.json\n"
        "S2\tquast\tfiles/missing.tsv\n"
    )
    parsed = report.read_manifest(manifest)
    assert parsed["fastp"]["S1"] == [tmp_path / "files" / "a.json"]
    assert parsed["citations"]["-"] == [data / "CITATIONS.md"]
    assert "unknown_tool" not in parsed
    assert "S2" not in parsed.get("quast", {})


def test_manifest_json_list_and_object(tmp_path, data):
    rows = [{"sample": "SCL2535", "tool": "mash", "file": str(data / "SCL2535.genome_size")}]
    for content in (rows, {"files": rows}):
        manifest = tmp_path / "manifest.json"
        manifest.write_text(json.dumps(content))
        assert report.read_manifest(manifest)["mash"]["SCL2535"] == [data / "SCL2535.genome_size"]


def test_manifest_empty_sample_is_run_level(tmp_path, data):
    manifest = tmp_path / "manifest.tsv"
    manifest.write_text(f"sample\ttool\tfile\n\tsample_status\t{data / 'sample_status.tsv'}\n")
    assert report.read_manifest(manifest)["sample_status"]["-"] == [data / "sample_status.tsv"]


# --------------------------------------------------------------------------------------------------
# Read QC
# --------------------------------------------------------------------------------------------------


def test_parse_fastp(data):
    result = report.parse_fastp(data / "SCL2535.fastp.json")
    assert result["reads_before"] == 2270906
    assert result["bases_before"] == 551085099
    assert result["bases_after"] == 472142971
    assert result["too_short_reads"] == 119844
    assert result["duplication_rate"] == pytest.approx(0.00348231)
    assert result["mean_length_before"] == pytest.approx(242.7, abs=0.05)


def test_parse_fastplong(data):
    result = report.parse_fastp(data / "ont.fastplong.json")
    assert result["reads_before"] == 1000
    assert result["bases_after"] == 33311690
    assert result["duplication_rate"] is None


def test_parse_mash_genome_size(data):
    assert report.parse_mash_genome_size(data / "SCL2535.genome_size") == 5012860


def test_parse_kraken2_species_rows_and_unclassified(data):
    species, unclassified = report.parse_kraken2(data / "SCL2333.kraken2.report.txt")
    assert unclassified == pytest.approx(0.24)
    assert list(species["name"]) == [
        "Klebsiella pneumoniae",
        "Raoultella ornithinolytica",
        "Klebsiella quasipneumoniae",
    ]
    assert species.loc[0, "percent"] == pytest.approx(27.88)
    assert species.loc[0, "reads"] == 329949


def test_parse_bracken_percent_of_classified_reads(data):
    table = report.parse_bracken(data / "SCL2333.bracken.tsv")
    assert table.loc[0, "name"] == "Klebsiella pneumoniae"
    assert table.loc[0, "percent"] == pytest.approx(96.143)
    assert table.loc[0, "reads"] == 1134253


# --------------------------------------------------------------------------------------------------
# Assembly
# --------------------------------------------------------------------------------------------------


def test_parse_quast(data):
    result = report.parse_quast(data / "511.1.quast.tsv")
    assert result["contigs_all"] == 397
    assert result["contigs"] == 48
    assert result["total_length"] == 4830321
    assert result["n50"] == 295713
    assert result["gc"] == pytest.approx(51.92)


def test_parse_filter_stats(data):
    result = report.parse_filter_stats(data / "511.1.filter_stats.tsv")
    assert result == {
        "assembly_contigs": 397,
        "assembly_length": 4907697,
        "assembly_n50": 295713,
        "annotation_contigs": 38,
        "annotation_length": 4823885,
        "annotation_n50": 295713,
    }


# --------------------------------------------------------------------------------------------------
# Typing
# --------------------------------------------------------------------------------------------------


def test_parse_mlst(data):
    result = report.parse_mlst(data / "100.1.mlst.tsv")
    assert result["scheme"] == "senterica_achtman_2"
    assert result["st"] == "32"
    assert result["alleles"].startswith("aroC(17) dnaN(18)")


def test_parse_mlst_without_scheme(tmp_path):
    path = tmp_path / "x.tsv"
    path.write_text("x.fasta\t-\t-\n")
    assert report.parse_mlst(path) == {"scheme": None, "st": None, "alleles": ""}


def test_parse_sistr(data):
    result = report.parse_sistr(data / "511.1.sistr.tab")
    assert result["serovar"] == "Agona"
    assert result["serogroup"] == "B"
    assert result["qc_status"] == "PASS"


def test_parse_kleborate_3_output(data):
    result = report.parse_kleborate(data / "SCL2333.kleborate.txt")
    assert result["species"] == "Klebsiella pneumoniae"
    assert result["ST"] == "ST25"
    assert result["K_locus"] == "KL2"
    assert result["Bla_ESBL_acquired"] == "CTX-M-15"
    assert result["virulence_score"] == "1"


def test_parse_kleborate_strips_module_prefixes(tmp_path):
    path = tmp_path / "klebsiella_pneumo_complex_output.txt"
    path.write_text(
        "strain\tenterobacterales__species__species\tklebsiella_pneumo_complex__mlst__ST\t"
        "klebsiella__abst__ST\n"
        "S1\tKlebsiella pneumoniae\tST258\t0\n"
    )
    result = report.parse_kleborate(path)
    assert result["species"] == "Klebsiella pneumoniae"
    assert result["ST"] == "ST258"


def test_parse_sccmec(data):
    result = report.parse_sccmec(data / "SCL1559.sccmec.tsv")
    assert (result["type"], result["subtype"], result["mecA"]) == ("I", "Ia", "+")


def test_sccmec_problem(data):
    # SCL1559: the target genes gave no type, so sccmec fell back to cassette coverage
    assert "cassette coverage only" in report.sccmec_problem(report.parse_sccmec(data / "SCL1559.sccmec.tsv"))
    assert report.sccmec_problem({"type": "II", "mecA": "+", "target_comment": ""}) is None
    assert report.sccmec_problem({"type": "-", "mecA": "-", "targets": "-"}) is None  # no SCCmec (MSSA)
    assert "no SCCmec type" in report.sccmec_problem({"type": "-", "mecA": "+", "targets": "mecA"})
    assert "several" in report.sccmec_problem({"type": "IV,V", "mecA": "+", "target_comment": ""})


# --------------------------------------------------------------------------------------------------
# Genome quality, AMR and mobile elements
# --------------------------------------------------------------------------------------------------


def test_parse_checkm2(data):
    result = report.parse_checkm2(data / "511.1.checkm2.tsv")
    assert result["completeness"] == 100.0
    assert result["contamination"] == pytest.approx(0.32)
    assert result["genome_size"] == 4823885
    assert result["model"].startswith("Neural Network")


def test_parse_amrfinderplus_4_columns(data):
    table = report.parse_amrfinderplus(data / "511.1.amrfinder.tsv")
    assert {"gene", "type", "subtype", "class", "identity", "coverage", "contig"} <= set(table.columns)
    amr = table[table["type"] == "AMR"]
    assert list(amr["gene"]) == ["mdsB", "mdsA", "fosA7.2"]
    assert amr["identity"].iloc[0] == pytest.approx(99.72)


def test_parse_amrfinderplus_3_columns(tmp_path):
    path = tmp_path / "old.tsv"
    path.write_text(
        "Gene symbol\tSequence name\tElement type\tElement subtype\tClass\tSubclass\t"
        "% Coverage of reference sequence\t% Identity to reference sequence\n"
        "blaKPC-2\tcarbapenem-hydrolyzing class A beta-lactamase KPC-2\tAMR\tAMR\tBETA-LACTAM\tCARBAPENEM\t100.00\t100.00\n"
    )
    table = report.parse_amrfinderplus(path)
    assert table.loc[0, "gene"] == "blaKPC-2"
    assert table.loc[0, "subclass"] == "CARBAPENEM"
    assert table.loc[0, "identity"] == 100.0


def test_parse_rgi(data):
    table = report.parse_rgi(data / "SCL1559.rgi.txt")
    assert list(table["gene"]) == ["norC", "arlR", "arlS", "ANT(9)-Ia", "ErmA"]
    assert list(table["cut_off"][:2]) == ["Strict", "Perfect"]
    assert table.loc[0, "identity"] == pytest.approx(70.64)
    # the ORF number is removed so contigs match the other tools
    assert list(table["contig"]) == ["contig_1", "contig_1", "contig_1", "contig_2", "contig_2"]


def test_parse_genomad(data):
    plasmids = report.parse_genomad(data / "100.1.genomad_plasmid.tsv")
    viruses = report.parse_genomad(data / "100.1.genomad_virus.tsv")
    assert len(plasmids) == 5
    assert plasmids.loc[0, "seq_name"] == "contig_26"
    assert len(viruses) == 2


def test_parse_mobsuite(data):
    contigs = report.parse_mobsuite_contigs(data / "SCL1559.mob_contigs.txt")
    plasmids = report.parse_mobsuite_mobtyper(data / "SCL1559.mobtyper.txt")
    assert (contigs["molecule_type"] == "plasmid").sum() == 7
    assert list(plasmids["primary_cluster_id"]) == ["AA849", "AA769", "AE018", "AE387"]
    assert plasmids.loc[0, "predicted_mobility"] == "mobilizable"


# --------------------------------------------------------------------------------------------------
# Methods
# --------------------------------------------------------------------------------------------------


def test_parse_citations_headings_and_list_items(data):
    entries = report.parse_citations(data / "CITATIONS.md")
    assert [entry["name"] for entry in entries] == ["nf-core", "Nextflow", "Autocycler", "GTDB-Tk", "LRGE"]
    assert entries[3]["url"] == "https://pubmed.ncbi.nlm.nih.gov/31730192/"
    assert entries[3]["text"].startswith("Chaumeil PA")
    assert entries[3]["text"].endswith("PMC7703759.")


def test_parse_software_versions_workflow_first(data):
    rows = report.parse_software_versions(data / "software_versions.yml")
    assert [(row["tool"], row["version"]) for row in rows[:2]] == [("gene2dis/mgap", "2.0.0"), ("Nextflow", "26.04.4")]
    rgi = next(row for row in rows if row["tool"] == "rgi")
    # one row per tool and version, listing every process that used it
    assert rgi["processes"] == ["GENE2DIS_MGAP:MGAP:RGI_LOAD", "GENE2DIS_MGAP:MGAP:RGI_MAIN"]


def test_parse_software_versions_keeps_versions_as_text(tmp_path):
    versions = tmp_path / "software_versions.yml"
    versions.write_text('"GENE2DIS_MGAP:MGAP:TOOL":\n    tool: 2.20\n    python: 3.10\n')
    rows = report.parse_software_versions(versions)
    # YAML would read 2.20 as the number 2.2
    assert {(row["tool"], row["version"]) for row in rows} == {("tool", "2.20"), ("python", "3.10")}


# --------------------------------------------------------------------------------------------------
# Run-level tables
# --------------------------------------------------------------------------------------------------


def test_parse_gtdbtk(data):
    result = report.parse_gtdbtk(data / "gtdbtk.bac120.summary.tsv")
    assert result["SCL2461"]["species"] == "Klebsiella pneumoniae"
    assert result["SCL2461"]["ani"] == pytest.approx(99.11)
    assert result["SCL2461"]["method"] == "ani_screen"
    # placed by the tree only: genus known, no species and no ANI
    assert result["CCG_02"]["species"] is None
    assert result["CCG_02"]["genus"] == "Micromonospora"
    assert result["CCG_02"]["ani"] is None


def test_parse_sample_status(data):
    table = report.parse_sample_status(data / "sample_status.tsv")
    assert list(table["status"]) == ["pass", "pass", "skipped"]
    low = table[table["sample"] == "SCL2535_low"].iloc[0]
    assert low["stage"] == "read_qc"
    assert low["reason"].startswith("16648045 bases after fastp")


# --------------------------------------------------------------------------------------------------
# Whole report from the fixtures
# --------------------------------------------------------------------------------------------------

FIXTURE_MANIFEST = [
    ("SCL2535", "fastp", "SCL2535.fastp.json"),
    ("SCL2535", "mash", "SCL2535.genome_size"),
    ("SCL2333", "kraken2", "SCL2333.kraken2.report.txt"),
    ("SCL2333", "bracken", "SCL2333.bracken.tsv"),
    ("SCL2333", "kleborate", "SCL2333.kleborate.txt"),
    ("511.1", "quast", "511.1.quast.tsv"),
    ("511.1", "filter_stats", "511.1.filter_stats.tsv"),
    ("511.1", "checkm2", "511.1.checkm2.tsv"),
    ("511.1", "sistr", "511.1.sistr.tab"),
    ("511.1", "amrfinderplus", "511.1.amrfinder.tsv"),
    ("100.1", "mlst", "100.1.mlst.tsv"),
    ("100.1", "genomad_plasmid", "100.1.genomad_plasmid.tsv"),
    ("100.1", "genomad_virus", "100.1.genomad_virus.tsv"),
    ("SCL1559", "rgi", "SCL1559.rgi.txt"),
    ("SCL1559", "sccmec", "SCL1559.sccmec.tsv"),
    ("SCL1559", "mobsuite_contigs", "SCL1559.mob_contigs.txt"),
    ("SCL1559", "mobsuite_mobtyper", "SCL1559.mobtyper.txt"),
    ("-", "gtdbtk", "gtdbtk.bac120.summary.tsv"),
    ("-", "sample_status", "sample_status.tsv"),
    ("-", "software_versions", "software_versions.yml"),
    ("-", "citations", "CITATIONS.md"),
]


def test_report_from_fixtures(tmp_path, data):
    manifest = tmp_path / "manifest.tsv"
    manifest.write_text(
        "sample\ttool\tfile\n" + "".join(f"{s}\t{tool}\t{data / name}\n" for s, tool, name in FIXTURE_MANIFEST)
    )
    html, tsv = tmp_path / "report.html", tmp_path / "summary.tsv"
    report.main(["--manifest", str(manifest), "--html", str(html), "--tsv", str(tsv)])
    page = html.read_text()
    for section in ("overview", "read-qc", "assembly", "quality", "typing", "amr", "mobile", "samples", "methods"):
        assert f'<section id="{section}"' in page
    # offline: Plotly is inlined, nothing is loaded from the network
    assert "<script src=" not in page
    summary = tsv.read_text().splitlines()
    assert summary[0].startswith("sample\t")
    samples = {line.split("\t")[0] for line in summary[1:]}
    assert {"SCL2535", "SCL2333", "511.1", "100.1", "SCL1559", "SCL2535_low", "SCL2461"} <= samples


def test_report_sections_need_their_inputs(tmp_path, data):
    manifest = tmp_path / "manifest.tsv"
    manifest.write_text(f"sample\ttool\tfile\n511.1\tquast\t{data / '511.1.quast.tsv'}\n")
    html = tmp_path / "report.html"
    report.main(["--manifest", str(manifest), "--html", str(html), "--tsv", str(tmp_path / "summary.tsv")])
    page = html.read_text()
    assert '<section id="assembly"' in page
    for section in ("read-qc", "quality", "typing", "amr", "mobile"):
        assert f'<section id="{section}"' not in page


def write_manifest(path, rows):
    path.write_text("sample\ttool\tfile\n" + "".join(f"{s}\t{tool}\t{file}\n" for s, tool, file in rows))
    return path


def test_samples_in_name_order(tmp_path, data):
    # sample_status.tsv is sorted as plain text (S1, S10, S2); the report sorts names naturally
    status = tmp_path / "sample_status.tsv"
    lines = (data / "sample_status.tsv").read_text().splitlines()
    status.write_text(lines[0] + "\n" + "".join(f"{name}\tillumina\tpass\n" for name in ("S1", "S10", "S2")))
    rows = [("-", "sample_status", status), ("S3", "quast", data / "511.1.quast.tsv")]
    run = report.RunData(report.read_manifest(write_manifest(tmp_path / "manifest.tsv", rows)), {})
    assert run.samples == ["S1", "S2", "S3", "S10"]


def test_sample_anchors_are_unique():
    names = ["A.1", "A_1", "A-1", "A_2e_1", "SCL2535_low"]
    slugs = [report.slug(name) for name in names]
    assert len(set(slugs)) == len(names)
    assert all(report.re.fullmatch(r"[A-Za-z0-9_-]+", slug) for slug in slugs)


def test_point_mutations_are_not_acquired_genes(tmp_path, data):
    lines = (data / "511.1.amrfinder.tsv").read_text().splitlines()
    fields = lines[-1].split("\t")
    # a gyrA point mutation: Element symbol, Type AMR, Subtype POINT
    fields[5], fields[8], fields[9] = "gyrA_D87Y", "AMR", "POINT"
    path = tmp_path / "511.1.amrfinder.tsv"
    path.write_text("\n".join(lines + ["\t".join(fields)]) + "\n")
    run = report.RunData(
        report.read_manifest(write_manifest(tmp_path / "manifest.tsv", [("511.1", "amrfinderplus", path)])), {}
    )
    assert run.amr_genes("511.1") == ["fosA7.2", "mdsA", "mdsB"]
    row = report.summary_rows(run).iloc[0]
    assert row["amrfinder_amr_gene_count"] == 3
    assert row["amrfinder_point_mutations"] == "gyrA_D87Y"


def test_heatmap_keeps_genes_with_the_same_short_label():
    # both labels shorten to "fluoroquinolone antibiotic · Escherichia coli sox…" (as in the Illumina run)
    genes = [f"Escherichia coli {name} with mutation conferring antibiotic resistance" for name in ("soxR", "soxS")]
    table = report.pd.DataFrame(
        {"gene": genes, "drug_group": "fluoroquinolone antibiotic", "identity": [99.0, 98.0], "coverage": 100.0}
    )
    page = report.amr_heatmap({"S1": table}, ["S1"], "drug_group")
    # Plotly writes the figure as JSON, with non-ASCII characters escaped
    rows = json.dumps([f"fluoroquinolone antibiotic · {gene}" for gene in genes], separators=(",", ":"))
    ticks = json.dumps(["fluoroquinolone antibiotic · Escherichia coli sox…"] * 2, separators=(",", ":"))
    assert f'"y":{rows}' in page
    assert f'"ticktext":{ticks}' in page


def test_rgi_without_strict_hits_says_so(tmp_path, data):
    lines = (data / "SCL1559.rgi.txt").read_text().splitlines()
    path = tmp_path / "SCL1559.rgi.txt"
    path.write_text(
        "\n".join(
            [lines[0]]
            + [line.replace("\tStrict\t", "\tLoose\t").replace("\tPerfect\t", "\tLoose\t") for line in lines[1:]]
        )
        + "\n"
    )
    run = report.RunData(
        report.read_manifest(write_manifest(tmp_path / "manifest.tsv", [("SCL1559", "rgi", path)])), {}
    )
    assert "No Perfect or Strict RGI hits were found." in report.section_amr(run).html()


def test_checkm2_groups_and_table_order(tmp_path, data):
    rows = [("511.1", "checkm2", data / "511.1.checkm2.tsv")]
    for sample, completeness, contamination in (("A", 85.0, 1.0), ("B", 99.0, 7.0), ("C", 80.0, 9.0)):
        path = tmp_path / f"{sample}.tsv"
        path.write_text(
            (data / "511.1.checkm2.tsv")
            .read_text()
            .replace("\t100.0\t0.32\t", f"\t{completeness}\t{contamination}\t")
            .replace("511.1", sample)
        )
        rows.append((sample, "checkm2", path))
    run = report.RunData(report.read_manifest(write_manifest(tmp_path / "manifest.tsv", rows)), {})
    groups = {s: run.checkm2_group(s) for s in ("511.1", "A", "B", "C")}
    assert groups == {"511.1": None, "A": "low_completeness", "B": "high_contamination", "C": "both"}
    page = report.section_quality(run).html()
    # genomes outside the limits come first in the table
    assert page.index('#sample-A"') < page.index("#sample-511_2e_1")


def test_typing_disagreement_first_and_no_second_call(tmp_path, data):
    mlst = tmp_path / "SCL2461.mlst.tsv"
    mlst.write_text("SCL2461.fasta\tecoli_achtman_4\t14464\tadk(1)\n")
    rows = [
        ("100.1", "mlst", data / "100.1.mlst.tsv"),
        ("SCL2461", "mlst", mlst),
        ("-", "gtdbtk", data / "gtdbtk.bac120.summary.tsv"),
    ]
    run_info = {"mlst_organisms": {"senterica_achtman_2": "Salmonella", "ecoli_achtman_4": "Escherichia"}}
    manifest = report.read_manifest(write_manifest(tmp_path / "manifest.tsv", rows))
    page = report.section_typing(report.RunData(manifest, run_info)).html()
    # GTDB-Tk calls SCL2461 Klebsiella: flagged and listed before 100.1
    assert page.index("#sample-SCL2461") < page.index("#sample-100_2e_1")
    assert "differs" in page
    # without GTDB-Tk or reads there is nothing to compare MLST with
    del manifest["gtdbtk"]
    page = report.section_typing(report.RunData(manifest, run_info)).html()
    assert "Neither GTDB-Tk nor a read classification ran" in page
    assert "Genus check" not in page


def test_species_check_sets_the_organism_and_typing_columns(tmp_path, data):
    """species_check.tsv: SCL2461 overridden to K. pneumoniae (mlst re-run: no ST, two phoE hits), 100.1 agrees."""
    mlst = tmp_path / "SCL2461.mlst.tsv"
    mlst.write_text("SCL2461.fasta\tklebsiella\t-\tgapA(3)\tphoE(1,1)\n")
    rows = [
        ("100.1", "mlst", data / "100.1.mlst.tsv"),
        ("SCL2461", "mlst", mlst),
        ("-", "gtdbtk", data / "gtdbtk.bac120.summary.tsv"),
        ("-", "species_check", data / "species_check.tsv"),
    ]
    run_info = {
        "mlst_organisms": {
            "senterica_achtman_2": "Salmonella",
            "ecoli_achtman_4": "Escherichia",
            "klebsiella": "Klebsiella_pneumoniae",
        }
    }
    run = report.RunData(report.read_manifest(write_manifest(tmp_path / "manifest.tsv", rows)), run_info)
    assert run.organism("SCL2461") == "Klebsiella_pneumoniae"
    assert run.organism("100.1") == "Salmonella"
    # routed as Klebsiella, as GTDB-Tk calls it: no genus disagreement left
    assert run.species_disagreement("SCL2461") is None
    assert any(
        note.startswith("Species check (override): MLST scheme ecoli_achtman_4") for note in run.warnings("SCL2461")
    )
    page = report.section_typing(run).html()
    assert "Species check" in page and "Closest reference (skani)" in page and "Species used" in page
    assert '<span class="flag">override</span>' in page
    # flagged first
    assert page.index("#sample-SCL2461") < page.index("#sample-100_2e_1")
    summary = report.summary_rows(run).set_index("sample")
    row = summary.loc["SCL2461"]
    # the MLST columns show the re-run, so the scheme's organism is the re-run scheme's
    assert (row["mlst_scheme"], row["mlst_st"], row["mlst_organism"]) == ("klebsiella", "", "Klebsiella_pneumoniae")
    assert (row["amrfinder_organism"], row["species_check_decision"]) == ("Klebsiella_pneumoniae", "override")
    assert (row["skani_species"], row["skani_ani"]) == ("Klebsiella pneumoniae", 99.11)


def test_build_manifest_finds_species_check(tmp_path):
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "accessory_scripts"))
    import build_report_manifest

    (tmp_path / "pipeline_info").mkdir()
    (tmp_path / "pipeline_info" / "species_check.tsv").write_text("sample\n")
    rows = build_report_manifest.manifest_rows(tmp_path, None)
    assert [row["tool"] for row in rows] == ["species_check"]
    # an mlst re-run replaces the automatic result, whatever the scheme's name sorts as
    for sample, rerun in (("S1", "vcholerae_2"), ("S2", None)):
        (tmp_path / sample / "annotation" / "mlst").mkdir(parents=True)
        (tmp_path / sample / "annotation" / "mlst" / f"{sample}.tsv").write_text("x\tecoli_achtman_4\t1\n")
        if rerun:
            (tmp_path / sample / "annotation" / "mlst" / f"{sample}.{rerun}.tsv").write_text(f"x\t{rerun}\t-\n")
    mlst = [
        Path(row["file"]).name for row in build_report_manifest.manifest_rows(tmp_path, None) if row["tool"] == "mlst"
    ]
    assert mlst == ["S1.vcholerae_2.tsv", "S2.tsv"]


def test_database_versions_in_overview(tmp_path, data):
    rows = [("-", "software_versions", data / "software_versions.yml")]
    run_info = {
        "databases": {"amrfinder_db": "/dbs/amrfinder", "bakta_db": "/dbs/bakta", "kraken2db": "/dbs/k2"},
        "database_versions": {"bakta_db": "6.0 full (2025-02-24)"},
    }
    run = report.RunData(report.read_manifest(write_manifest(tmp_path / "manifest.tsv", rows)), run_info)
    page = report.section_overview(run).html()
    assert "2025-12-03.1" in page  # from software_versions.yml
    assert "6.0 full (2025-02-24)" in page  # from run-info
    assert "not recorded" in page  # Kraken2 databases carry no version


def test_build_manifest_database_versions(tmp_path):
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "accessory_scripts"))
    import build_report_manifest

    (tmp_path / "bakta").mkdir()
    (tmp_path / "bakta" / "version.json").write_text('{"date": "2025-02-24", "major": 6, "minor": 0, "type": "full"}')
    (tmp_path / "gtdb" / "metadata").mkdir(parents=True)
    (tmp_path / "gtdb" / "metadata" / "metadata.txt").write_text("VERSION_DATA=r232\n")
    (tmp_path / "rgi").mkdir()
    (tmp_path / "rgi" / "card_database_v4.0.1_all.fasta").write_text("")
    (tmp_path / "k2").mkdir()
    databases = {name: str(tmp_path / name) for name in ("bakta", "gtdb", "rgi", "k2")}
    databases = dict(zip(("bakta_db", "gtdbtk_db", "rgi_db", "kraken2db"), databases.values()))
    assert build_report_manifest.database_versions(databases) == {
        "bakta_db": "6.0 full (2025-02-24)",
        "gtdbtk_db": "r232",
        "rgi_db": "CARD 4.0.1",
    }


def test_build_manifest_covers_every_sample_tool():
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "accessory_scripts"))
    import build_report_manifest

    # rebuilding the report from a results folder finds every per-sample input the report reads
    assert set(report.SAMPLE_TOOLS) - set(build_report_manifest.SAMPLE_PATTERNS) == set()


def test_build_manifest_finds_plasmid_consensus(tmp_path):
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "accessory_scripts"))
    import build_report_manifest

    folder = tmp_path / "SCL2535" / "annotation" / "plasmid_consensus"
    folder.mkdir(parents=True)
    (folder / "SCL2535.plasmid_consensus.tsv").write_text("sample\tcontig\n")
    rows = build_report_manifest.manifest_rows(tmp_path, None)
    assert [(row["sample"], row["tool"]) for row in rows] == [("SCL2535", "plasmid_consensus")]


# --------------------------------------------------------------------------------------------------
# Large runs, sample cards and the workbook
# --------------------------------------------------------------------------------------------------


def test_outliers_tukey_fences_and_limits():
    values = {f"S{i}": 10.0 + i % 3 for i in range(20)} | {"high": 40.0, "low_limit": 9.0}
    assert report.outliers(values) == {"high"}
    assert report.outliers(values, low=9.5) == {"high", "low_limit"}
    # no spread: only the limits flag anything
    assert report.outliers({"A": 1.0, "B": 1.0}) == set()


def test_outlier_labels_group_close_points():
    values = {"a": 0.0, "b": 0.1, "c": 50.0, "d": 100.0}
    labels = report.outlier_labels(values, ["a", "b", "d"])
    assert labels == {"a": "a +1", "d": "d"}


def test_large_run_assembly_is_a_distribution(tmp_path, data):
    rows = [(f"S{i:03d}", "quast", data / "511.1.quast.tsv") for i in range(report.LARGE_RUN + 1)]
    run = report.RunData(report.read_manifest(write_manifest(tmp_path / "manifest.tsv", rows)), {})
    page = report.section_assembly(run).html()
    assert '"type":"box"' in page
    # no chart with one row per sample (bar charts put the sample names on the y axis)
    assert '"y":["S' not in page
    # the per-sample table is still there, in a scrolling box
    assert "S100" in page and f"max-height:{report.TABLE_SCROLL_HEIGHT}px" in page
    small = report.RunData(report.read_manifest(write_manifest(tmp_path / "small.tsv", rows[:3])), {})
    assert '"y":["S' in report.section_assembly(small).html()


def test_heatmap_empty_cells_have_no_hover_text(tmp_path, data):
    other = tmp_path / "other.amrfinder.tsv"
    lines = (data / "511.1.amrfinder.tsv").read_text().splitlines()
    other.write_text("\n".join(lines[:2]) + "\n")  # header and one hit
    rows = [("511.1", "amrfinderplus", data / "511.1.amrfinder.tsv"), ("other", "amrfinderplus", other)]
    run = report.RunData(report.read_manifest(write_manifest(tmp_path / "manifest.tsv", rows)), {})
    page = report.section_amr(run).html()
    assert '"hoverongaps":false' in page
    assert "not found" not in page


def test_cards_collapsed_with_flagged_samples_first(tmp_path, data):
    manifest = write_manifest(tmp_path / "manifest.tsv", [(s, tool, data / name) for s, tool, name in FIXTURE_MANIFEST])
    html = tmp_path / "report.html"
    report.main(["--manifest", str(manifest), "--html", str(html), "--tsv", str(tmp_path / "summary.tsv")])
    page = html.read_text()
    assert '<details class="card" id="sample-SCL2535_5f_low"' in page
    assert 'id="sample-search"' in page
    listing = page[page.index('<ul class="sample-list">') :]
    chips = listing[: listing.index("</ul>")].split("<li ")[1:]
    flagged = ["flagged" in chip for chip in chips]
    # every flagged sample comes before every other one, and the skipped sample is flagged
    assert flagged == sorted(flagged, reverse=True) and any(flagged) and not all(flagged)
    assert 'class="sample-chip flagged" href="#sample-SCL2535_5f_low"' in listing


def test_xlsx_summary_and_one_sheet_per_tool(tmp_path, data):
    pytest.importorskip("openpyxl", minversion="3.1.0")  # the oldest pandas 2.2 accepts
    manifest = write_manifest(tmp_path / "manifest.tsv", [(s, tool, data / name) for s, tool, name in FIXTURE_MANIFEST])
    xlsx = tmp_path / "summary.xlsx"
    report.main(
        [
            "--manifest",
            str(manifest),
            "--html",
            str(tmp_path / "r.html"),
            "--tsv",
            str(tmp_path / "s.tsv"),
            "--xlsx",
            str(xlsx),
        ]
    )
    sheets = report.pd.read_excel(xlsx, sheet_name=None)
    assert list(sheets)[0] == "Summary"
    for name in ("Sample status", "Reads", "QUAST", "CheckM2", "MLST", "GTDB-Tk", "Kleborate", "SCCmec", "RGI"):
        assert name in sheets
    assert "Predicted_DNA" not in sheets["RGI"]
    assert set(sheets["SCCmec"]["sample"]) == {"SCL1559"}
    assert len(sheets["Summary"]) == len((tmp_path / "s.tsv").read_text().splitlines()) - 1


def test_xlsx_text_columns_cleaned(tmp_path):
    openpyxl = pytest.importorskip("openpyxl", minversion="3.1.0")
    summary = report.pd.read_csv(
        io.StringIO("sample\tnote\nS1\tok\x01bad\nS2\t" + "x" * 40000 + "\n"), sep="\t", dtype=str
    )
    xlsx = tmp_path / "summary.xlsx"
    report.write_xlsx(xlsx, summary, {})
    sheet = openpyxl.load_workbook(xlsx)["Summary"]
    assert sheet["B2"].value == "okbad"
    assert len(sheet["B3"].value) == report.XLSX_CELL


def test_plasmid_consensus_section_and_summary(tmp_path, data):
    manifest = write_manifest(
        tmp_path / "manifest.tsv",
        [("SCL2535", "plasmid_consensus", data / "plasmid" / "SCL2535.plasmid_consensus.tsv")],
    )
    html, tsv, xlsx = tmp_path / "report.html", tmp_path / "summary.tsv", tmp_path / "summary.xlsx"
    report.main(["--manifest", str(manifest), "--html", str(html), "--tsv", str(tsv), "--xlsx", str(xlsx)])
    page = html.read_text()
    assert '<section id="plasmids"' in page
    assert "Contigs with plasmid evidence" in page
    assert "10 plasmid, 8 candidate" in page  # sample card
    header, row = tsv.read_text().splitlines()
    summary = dict(zip(header.split("\t"), row.split("\t")))
    assert summary["plasmid_consensus_plasmid"] == "10"
    assert summary["plasmid_consensus_candidate"] == "8"
    assert summary["plasmid_consensus_check"] == "0"
    assert summary["plasmid_replicons"].split(", ")[:3] == ["IncFIB(pKPHS1)", "IncFIB(K)", "IncFII(K)"]
    import openpyxl

    sheet = openpyxl.load_workbook(xlsx)["Plasmid consensus"]
    assert sheet.max_row == 41  # header + 40 contigs
    assert [cell.value for cell in sheet[1]][:3] == ["sample", "contig", "description"]


def test_plasmid_contig_table_names_each_replicon_once(tmp_path, data):
    # PlasmidFinder can hit one replicon several times on a contig (ONT SCL10571: Col440I twice)
    text = (data / "plasmid" / "SCL2535.plasmid_consensus.tsv").read_text()
    edited = tmp_path / "consensus.tsv"
    edited.write_text(text.replace("\tIncFIB(pKPHS1)\t97.32", "\tIncFIB(pKPHS1),IncFIB(pKPHS1)\t97.32,96.10"))
    assert edited.read_text() != text
    manifest = write_manifest(tmp_path / "manifest.tsv", [("SCL2535", "plasmid_consensus", edited)])
    html = tmp_path / "report.html"
    report.main(["--manifest", str(manifest), "--html", str(html), "--tsv", str(tmp_path / "summary.tsv")])
    page = html.read_text()
    assert "IncFIB(pKPHS1)" in page
    assert "IncFIB(pKPHS1), IncFIB(pKPHS1)" not in page


# --------------------------------------------------------------------------------------------------
# ISEScan
# --------------------------------------------------------------------------------------------------


def test_parse_isescan(data):
    table = report.parse_isescan(data / "SCL2535.isescan.tsv")
    assert len(table) == 31
    assert table["isLen"].sum() == 38560
    assert report.is_complete(table) == 13
    families = report.is_families(table)
    # most elements first, ties in natural name order
    assert list(families.items())[:4] == [("IS3", 10), ("ISNCY", 6), ("IS1", 3), ("IS6", 2)]
    assert report.is_family_text(families, 2) == "IS3 10, ISNCY 6, ..."
    assert report.parse_isescan(data / "empty.isescan.tsv").empty


def test_isescan_section_summary_and_sheet(tmp_path, data):
    manifest = write_manifest(
        tmp_path / "manifest.tsv",
        [
            ("SCL2535", "isescan", data / "SCL2535.isescan.tsv"),
            ("SCL2535", "plasmid_consensus", data / "plasmid" / "SCL2535.plasmid_consensus.tsv"),
            ("SCL2333", "isescan", data / "empty.isescan.tsv"),
        ],
    )
    html, tsv, xlsx = tmp_path / "report.html", tmp_path / "summary.tsv", tmp_path / "summary.xlsx"
    report.main(["--manifest", str(manifest), "--html", str(html), "--tsv", str(tsv), "--xlsx", str(xlsx)])
    page = html.read_text()
    assert '<section id="isescan"' in page
    assert "31 (13 complete): IS3 10, ISNCY 6, IS1 3, ..." in page  # sample card
    assert "none found" in page  # the sample whose table has no element
    assert "Other families" in page  # 11 families: 7 named in the chart
    assert "On plasmid contigs" in page
    assert "On candidate plasmid contigs" in page
    lines = tsv.read_text().splitlines()
    rows = {line.split("\t")[0]: dict(zip(lines[0].split("\t"), line.split("\t"))) for line in lines[1:]}
    assert rows["SCL2535"]["isescan_elements"] == "31"
    assert rows["SCL2535"]["isescan_complete"] == "13"
    assert rows["SCL2535"]["isescan_families"].startswith("IS3:10,ISNCY:6,IS1:3,")
    assert rows["SCL2333"]["isescan_elements"] == "0"
    assert rows["SCL2333"]["isescan_families"] == ""
    import openpyxl

    sheet = openpyxl.load_workbook(xlsx)["ISEScan"]
    assert sheet.max_row == 32  # header + 31 elements of SCL2535
    assert [cell.value for cell in sheet[1]][:3] == ["sample", "seqID", "family"]


def test_isescan_elements_on_plasmid_contigs(tmp_path, data):
    run = report.RunData(
        {
            "isescan": {"SCL2535": [data / "SCL2535.isescan.tsv"]},
            "plasmid_consensus": {"SCL2535": [data / "plasmid" / "SCL2535.plasmid_consensus.tsv"]},
        },
        {},
    )
    # contigs 15, 20, 21, 25 and 31 are called plasmid and carry 13 of the 31 elements
    assert report.is_on_plasmids(run, "SCL2535") == 13
    # candidates (one tool only): contig_22 with 3 elements, contig_29 and contig_38 with 1 each
    assert report.is_on_plasmids(run, "SCL2535", "candidate") == 5
    without = report.RunData({"isescan": {"SCL2535": [data / "SCL2535.isescan.tsv"]}}, {})
    assert report.is_on_plasmids(without, "SCL2535") is None
    # a consensus table without a contig column (hand-built or older) is skipped, not an error
    broken = tmp_path / "consensus.tsv"
    broken.write_text("consensus\nplasmid\n")
    no_contig = report.RunData(
        {"isescan": {"SCL2535": [data / "SCL2535.isescan.tsv"]}, "plasmid_consensus": {"SCL2535": [broken]}}, {}
    )
    assert report.is_on_plasmids(no_contig, "SCL2535") is None


def test_isescan_large_run_uses_distributions(tmp_path, data):
    rows = [(f"S{i}", "isescan", data / "SCL2535.isescan.tsv") for i in range(report.LARGE_RUN + 1)]
    manifest = write_manifest(tmp_path / "manifest.tsv", rows)
    html = tmp_path / "report.html"
    report.main(["--manifest", str(manifest), "--html", str(html), "--tsv", str(tmp_path / "summary.tsv")])
    page = html.read_text()
    section = page[page.index('<section id="isescan"') :]
    section = section[: section.index("</section>")]
    assert "Complete IS elements" in section
    assert "Other families" not in section
